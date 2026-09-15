# 语音祝福响度统一（安全版）：
# 从群收集平台的原始链接重新下载全部录音，测响度后按差值纯增益提升，再用限幅器兜底（绝不削波）。
# 目标 -10 LUFS；用法: python scripts/boost-voices.py [祝福文件夹路径]
import glob, os, subprocess, sys, json
import openpyxl
import imageio_ffmpeg

FFMPEG = imageio_ffmpeg.get_ffmpeg_exe()
OUT_AUDIO = "public/welcome/voices"
TARGET_I = -10.0
LIMIT = 0.794  # ≈ -2 dBTP

def fetch(url, dst):
    r = subprocess.run(["curl", "-sL", "--max-time", "120", url, "-o", dst], capture_output=True)
    return r.returncode == 0 and os.path.getsize(dst) > 1000

def measure(f):
    r = subprocess.run([FFMPEG, "-hide_banner", "-i", f,
                        "-af", f"loudnorm=I={TARGET_I}:print_format=json",
                        "-f", "null", "-"], capture_output=True, text=True)
    stderr = r.stderr
    js = stderr[stderr.rfind("{"):stderr.rfind("}") + 1]
    return json.loads(js)

def volume_peak(f):
    r = subprocess.run([FFMPEG, "-hide_banner", "-i", f, "-af", "volumedetect",
                        "-f", "null", "-"], capture_output=True, text=True)
    mean = peak = None
    for line in r.stderr.splitlines():
        if "mean_volume" in line: mean = float(line.split(":")[1].replace("dB", "").strip())
        if "max_volume" in line: peak = float(line.split(":")[1].replace("dB", "").strip())
    return mean, peak

def main():
    candidates = glob.glob("祝福*") or glob.glob("*/祝福*") or ["祝福"]
    folder = candidates[0]
    path = os.path.join(folder, "新生礼物资料收集_群收集.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    tmp = os.path.join(folder, ".tmp-audio")

    print(f"处理 {ws.max_row - 1} 条录音（目标 {TARGET_I} LUFS，限幅 {LIMIT}）")
    results = []
    for row in range(2, ws.max_row + 1):
        seq = ws.cell(row=row, column=1).value
        if seq is None: continue
        nn = f"{int(seq):02d}"
        name = ws.cell(row=row, column=4).value
        audio_cell = ws.cell(row=row, column=6)
        url = audio_cell.hyperlink.target if audio_cell.hyperlink else None
        raw = os.path.join(folder, f".raw-{nn}")
        if not url or not fetch(url, raw):
            print(f"  seq{nn} {name}: 下载失败 ✗"); continue

        m = measure(raw)
        gain = round(TARGET_I - float(m["input_i"]), 2)
        out = os.path.join(OUT_AUDIO, f"seq{nn}.m4a")
        subprocess.run([FFMPEG, "-y", "-hide_banner", "-loglevel", "error", "-i", raw,
                        "-af", f"volume={gain}dB,alimiter=limit={LIMIT}:level=disabled",
                        "-c:a", "aac", "-b:a", "96k", out], check=True)
        mean, peak = volume_peak(out)
        results.append((nn, name, gain, mean, peak))
        print(f"  seq{nn} {name}: 增益 {gain:+.1f}dB → mean {mean:.1f}dB peak {peak:.1f}dB ✓")
        os.remove(raw)

    print("完成")
    clipped = [r for r in results if r[4] > -0.3]
    if clipped:
        print("警告：以下文件仍接近满幅，请检查：", clipped)

if __name__ == "__main__":
    main()
